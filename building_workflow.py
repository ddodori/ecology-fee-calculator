"""건축물대장 PNU 자동생성과 토지조서 대조를 한 번에 수행하는 엔진.

기존 pnu_core.process와 match_core.match_to_ledger의 검증 로직을 재사용합니다.
"""
from __future__ import annotations
from dataclasses import dataclass, field
from typing import Optional
from io import BytesIO
from openpyxl import load_workbook

from pnu_core import process
from match_core import match_to_ledger


@dataclass
class BuildingWorkflowResult:
    pnu_bytes: bytes
    pnu_stats: dict
    source_kind: str
    match_bytes: Optional[bytes] = None
    match_stats: dict = field(default_factory=dict)
    ledger_sheet: str = ""


def run_building_workflow(
    building_bytes: bytes,
    ledger_bytes: bytes | None = None,
    code_bytes: bytes | None = None,
) -> BuildingWorkflowResult:
    """PNU 생성은 1회만 실행하고, 결과의 중복제거 목록을 바로 대조에 사용."""
    # 기존 V1.2 생성 결과(PNU_중복제거 포함)도 그대로 사용할 수 있도록 지원합니다.
    try:
        book = load_workbook(BytesIO(building_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("건축물대장 엑셀을 읽지 못했습니다. .xlsx 파일인지 확인해 주세요.") from exc
    try:
        is_processed = "PNU_중복제거" in book.sheetnames
        if is_processed:
            pnu_bytes, source_kind = building_bytes, "기존 PNU 생성결과"
            pnu_stats = {}
            if "처리요약" in book.sheetnames:
                for row in book["처리요약"].iter_rows(min_row=2, max_col=2, values_only=True):
                    if row[0] in ("전체 데이터", "정상 PNU 행", "고유 PNU 수", "제외 행", "중복 PNU 행 수"):
                        stat_name = {
                            "전체 데이터": "전체", "정상 PNU 행": "정상",
                            "고유 PNU 수": "고유 PNU", "제외 행": "제외",
                            "중복 PNU 행 수": "중복 PNU 행 수",
                        }[row[0]]
                        pnu_stats[stat_name] = int(row[1] or 0)
        else:
            pnu_bytes, pnu_stats, source_kind = process(building_bytes, code_bytes)
    finally:
        book.close()
    result = BuildingWorkflowResult(pnu_bytes=pnu_bytes, pnu_stats=pnu_stats,
                                    source_kind=source_kind)
    if ledger_bytes is not None:
        match_bytes, match_stats, ledger_sheet, _ = match_to_ledger(
            ledger_bytes, pnu_bytes, code_bytes
        )
        result.match_bytes = match_bytes
        result.match_stats = match_stats
        result.ledger_sheet = ledger_sheet
    return result
