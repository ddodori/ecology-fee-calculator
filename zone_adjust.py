"""2-2 용도지역별 편입면적 비례 보정 엔진 (V1.4.12).

교차분석 결과를 이어받아 공간 경고를 기록하고, 각 PNU에 속한 피처에
필지 편입면적을 비례배분합니다. 보정금액/법정 부담금 계산은 하지 않습니다.

반올림은 최대잔여법(largest remainder)으로 분배합니다. 각 피처의 정확한
비례 몫을 구한 뒤 필지별 원본 편입면적의 유효 소수 자릿수에 따라
1m² / 0.1m² / 0.01m² 단위로 내림하고 남은 단위를 소수 잔여가
큰 피처부터 배분합니다. 이때 PNU별 합계는 원본 편입면적과 일치합니다.
예: 편입면적 1,000㎡인 필지는 최종 용도지역별 면적도 정수㎡입니다.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_FLOOR
from io import BytesIO

import pandas as pd


@dataclass
class AdjustmentResult:
    excel_bytes: bytes
    adjusted_parcels: int
    pending_parcels: int
    total_parcels: int
    adjusted_area_m2: float
    pending_area_m2: float
    summary: pd.DataFrame
    preview: pd.DataFrame
    pending: pd.DataFrame
    review: pd.DataFrame
    review_parcels: int
    validation: pd.DataFrame
    detail: pd.DataFrame
    excluded_zero_segments: int
    excluded_zero_zones: int


def _as_decimal(value) -> Decimal | None:
    if value is None or pd.isna(value) or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).replace(',', '').strip())
        return number if number.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def _exact_units(area: Decimal, precision: int) -> int | None:
    step = Decimal(1).scaleb(-precision)
    units = area / step
    rounded = units.to_integral_value()
    # Uploaded CSV/SHP floating point noise tolerance: 1e-7 of a chosen unit.
    if abs(units - rounded) > Decimal('0.0000001'):
        return None
    return int(rounded)


def _parcel_precision(area: Decimal, max_precision: int = 2) -> int | None:
    """Smallest 1 / 0.1 / 0.01㎡ unit that can exactly represent a parcel target.

    Treat 1000, 1000.0 and 1000.00 identically as *integer* source areas.
    _exact_units provides a tiny tolerance for numeric serialization noise.
    """
    return next((digits for digits in range(max_precision + 1)
                 if _exact_units(area, digits) is not None), None)


def _apportion(areas: list[Decimal], target_units: int) -> list[int]:
    """Hamilton allocation: nonnegative exact unit totals, stable tie breaking."""
    if target_units < 0 or any(a < 0 for a in areas):
        raise ValueError('음수 면적은 보정할 수 없습니다.')
    total = sum(areas)
    if total <= 0:
        raise ValueError('교차 구적면적 합계가 0 이하입니다.')
    quotas = [Decimal(target_units) * a / total for a in areas]
    bases = [int(q.to_integral_value(rounding=ROUND_FLOOR)) for q in quotas]
    missing = target_units - sum(bases)
    if missing < 0 or missing > len(bases):
        raise ArithmeticError('잔여 단위 배분량이 비정상적입니다.')
    rank = sorted(range(len(areas)), key=lambda i: (-(quotas[i] - bases[i]), -areas[i], i))
    for i in rank[:missing]:
        bases[i] += 1
    assert sum(bases) == target_units
    return bases


def _write_xlsx(sheets: dict[str, pd.DataFrame], precision: int) -> bytes:
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        from openpyxl.styles import Alignment, Font, PatternFill
        for name, table in sheets.items():
            table.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='DBEAFE')
                cell.font = Font(name='맑은 고딕', bold=True, color='1E3A5F')
                cell.alignment = Alignment(vertical='center')
            ws.row_dimensions[1].height = 24
            for j, col_name in enumerate(table.columns, start=1):
                column_letter = ws.cell(1, j).column_letter
                sample_width = max((len(str(x)) for x in table[col_name].head(60) if pd.notna(x)), default=8)
                ws.column_dimensions[column_letter].width = max(14, min(34, max(len(str(col_name)), sample_width) + 3))
                if 'PNU' in str(col_name) or '코드' in str(col_name) or 'MNUM' in str(col_name):
                    for cells in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                        for c in cells:
                            c.number_format = '@'
                # Avoid '#,##0.##': Excel can render integers as '402.' with
                # the trailing dot. Integer amounts show no dot at all.
                if '면적' in str(col_name) or str(col_name) in ('필지편입면적', '필지대장면적') or '차이_m2' in str(col_name):
                    for cells in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                        for c in cells:
                            if isinstance(c.value, (float, int)) and not isinstance(c.value, bool):
                                if '교차구적' in str(col_name):
                                    c.number_format = '#,##0.00'
                                else:
                                    dec = Decimal(str(c.value))
                                    if abs(dec - dec.to_integral_value()) < Decimal('0.000000001'):
                                        c.number_format = '#,##0'
                                    elif abs(dec*10 - (dec*10).to_integral_value()) < Decimal('0.000000001'):
                                        c.number_format = '#,##0.0'
                                    else:
                                        # Keep even invalid (> 0.01㎡) source precision
                                        # visible in audit/pending sheets: 30.123 must
                                        # not display misleadingly as 30.12.
                                        places = max(2, -dec.normalize().as_tuple().exponent)
                                        c.number_format = '#,##0.' + ('0' * min(places, 10))
    return stream.getvalue()


def adjust_zone_areas(detail: pd.DataFrame, parcel_validation: pd.DataFrame,
                      precision: int = 2) -> AdjustmentResult:
    """Return an XLSX and audit frames. Does not overwrite the uploaded source.

    Required input columns of ``detail`` and ``parcel_validation`` are produced
    by ``analyze_zone_overlay``. Keep all intersection rows for apportionment.
    """
    if precision != 2:
        raise ValueError('V1.4.12은 최대 0.01㎡까지 필지별 자동 보정만 지원합니다.')
    required_detail = {'PNU','지목','원본편입면적','원본대장면적',
                       '용도지역코드','용도지역명','코드검증','교차구적면적_m2','분석행ID'}
    required_validation = {'PNU','분석대상면적_m2','원본편입면적','원본대장면적',
                           '필지원본행수','공간검증'}
    if not required_detail.issubset(detail.columns) or not required_validation.issubset(parcel_validation.columns):
        raise ValueError('V1.4.6 교차분석 결과가 필요합니다. ① 교차분석부터 다시 실행하세요.')

    raw = detail.copy().reset_index(drop=True)
    raw['최종보정면적_m2'] = float('nan')
    raw['보정상태'] = '보정 보류'
    raw['보류사유'] = '분석자료 확인 필요'
    raw['보정차이_m2'] = float('nan')
    raw['적용보정단위_m2'] = float('nan')
    validations = parcel_validation.copy().reset_index(drop=True)
    validations['최종보정합계_m2'] = float('nan')
    validations['편입면적차이_m2'] = float('nan')
    validations['보정상태'] = '보정 보류'
    validations['보류사유'] = ''
    validations['공간검토사유'] = ''
    validations['적용보정단위_m2'] = float('nan')

    # Spatial diagnostics are checked at PNU-level before any forced normalization.
    for idx, row in validations.iterrows():
        pnu = str(row['PNU']).strip()
        locs = raw.index[raw['PNU'].astype(str).str.strip() == pnu].tolist()
        why = []
        if len(pnu) != 19 or not pnu.isdigit():
            why.append('PNU 형식 오류')
        # Spatial coverage warnings are diagnostics, not automatic blockers.
        # Ordinary geometry/ledger area discrepancies can be allocated by ratio.
        # Coverage warnings remain visible in the output for human review.
        geo_status = str(row['공간검증'])
        needs_geo_review = geo_status not in ('정상', 'PNU 확인 필요')
        if needs_geo_review:
            validations.at[idx, '공간검토사유'] = geo_status
        if int(row['필지원본행수']) != 1:
            why.append('동일 PNU가 지적도 원본에 여러 행 존재')
        incl = _as_decimal(row['원본편입면적'])
        cadastral = _as_decimal(row['원본대장면적'])
        if incl is None or incl < 0:
            why.append('편입면적 누락 또는 오류')
        if cadastral is None or cadastral <= 0:
            why.append('대장면적 누락 또는 오류')
        if incl is not None and cadastral is not None and incl > cadastral + Decimal('0.0000001'):
            why.append('편입면적이 대장면적 초과')
        # Allocation resolution is PER PNU, not global. 1000.0 -> 1㎡,
        # 1000.3 -> 0.1㎡, and 1000.35 -> 0.01㎡.
        target_precision = (_parcel_precision(incl, precision)
                            if incl is not None and incl >= 0 else None)
        target_units = (_exact_units(incl, target_precision)
                        if target_precision is not None else None)
        if target_units is None and incl is not None and incl >= 0:
            why.append('편입면적에 0.01㎡ 미만 소수 자릿수 존재 (원본 확인 필요)')
        if not locs:
            why.append('용도지역 교차 결과 없음')
        areas = []
        for r in locs:
            a = _as_decimal(raw.at[r, '교차구적면적_m2'])
            if a is None or a <= 0:
                why.append('교차면적 누락 또는 0 이하')
            areas.append(a)
            if str(raw.at[r, '코드검증']) != '정상':
                why.append('용도지역코드 확인 필요')
        # Integrity: all split sections of a PNU must reference the same source areas.
        for r in locs:
            if _as_decimal(raw.at[r,'원본편입면적']) != incl or _as_decimal(raw.at[r,'원본대장면적']) != cadastral:
                why.append('원본 필지 면적값 상충')
                break

        if why:
            reason = ' / '.join(dict.fromkeys(why))
            validations.at[idx, '보류사유'] = reason
            raw.loc[locs, '보류사유'] = reason
            continue
        amounts = _apportion(areas, target_units)
        units_multiplier = 10 ** target_precision
        applied_step = 1 / units_multiplier
        for pos, area_units in zip(locs, amounts):
            adjusted = area_units / units_multiplier
            raw.at[pos,'최종보정면적_m2'] = adjusted
            raw.at[pos,'적용보정단위_m2'] = applied_step
            raw.at[pos,'보정차이_m2'] = adjusted - float(raw.at[pos,'교차구적면적_m2'])
            raw.at[pos,'보정상태'] = '보정 완료·공간검토필요' if needs_geo_review else '정상 보정'
            raw.at[pos,'보류사유'] = ''
        sum_units = sum(amounts)
        if sum_units != target_units:
            raise ArithmeticError(f'{pnu}: 보정면적 합이 편입면적과 다릅니다.')
        validations.at[idx,'최종보정합계_m2'] = sum_units / units_multiplier
        validations.at[idx,'적용보정단위_m2'] = applied_step
        validations.at[idx,'편입면적차이_m2'] = (sum_units-target_units)/units_multiplier
        validations.at[idx,'보정상태'] = '보정 완료·공간검토필요' if needs_geo_review else '정상 보정'
        validations.at[idx,'보류사유'] = ''

    # A successfully adjusted intersection can receive 0㎡ when the PNU has
    # smaller geometry pieces than the target precision (1 / 0.1 / 0.01㎡).
    # Preserve these rows for audit, but remove them from final outputs.
    completed_statuses = ['정상 보정', '보정 완료·공간검토필요']
    zero_segment_mask = (raw['보정상태'].isin(completed_statuses)
                         & pd.to_numeric(raw['최종보정면적_m2'], errors='coerce').eq(0))
    excluded_zero_segments = raw.loc[zero_segment_mask].copy()
    clean = raw[raw['보정상태'].isin(completed_statuses)].copy()
    # The main output uses one row per PNU + zone code + zone name, not each
    # geometric fragment. Carry all original parcel columns through the groupby.
    # The earlier V1.4.6 had dropped columns such as 법정동, 지번, 건축물 여부.
    derived = {'분석행ID', 'PNU', '지목', '원본편입면적', '원본대장면적',
               'MNUM', '용도지역코드', '용도지역명', '코드검증', '교차구적면적_m2',
               '최종보정면적_m2', '보정상태', '보류사유', '보정차이_m2'}
    source_extra_fields = [col for col in detail.columns if col not in derived]
    output_cols = (['PNU'] + source_extra_fields + ['지목', '용도지역코드', '용도지역명',
                   '필지편입면적', '필지대장면적', '교차구적면적_m2', '최종보정면적_m2', '보정상태'])
    if not clean.empty:
        clean['필지편입면적'] = clean['원본편입면적']
        clean['필지대장면적'] = clean['원본대장면적']
        aggregations = {
            '지목': ('지목', 'first'),
            '필지편입면적': ('필지편입면적', 'first'),
            '필지대장면적': ('필지대장면적', 'first'),
            '교차구적면적_m2': ('교차구적면적_m2', 'sum'),
            '최종보정면적_m2': ('최종보정면적_m2', 'sum'),
            '보정상태': ('보정상태', lambda x: '보정 완료·공간검토필요' if (x == '보정 완료·공간검토필요').any() else '정상 보정'),
        }
        for col in source_extra_fields:
            aggregations[col] = (col, 'first')
        clean = (clean.groupby(['PNU', '용도지역코드', '용도지역명'], as_index=False, dropna=False)
                .agg(**aggregations)
                .sort_values(['PNU', '용도지역코드'], kind='stable'))
        clean['최종보정면적_m2'] = clean['최종보정면적_m2'].round(precision)
        # A group can contain several fragments; filter AFTER aggregating so
        # only actually-zero PNU+zone rows disappear from the primary Excel.
        excluded_zero_zones = int(clean['최종보정면적_m2'].eq(0).sum())
        clean = clean.loc[clean['최종보정면적_m2'].gt(0), output_cols].copy()
    else:
        clean = pd.DataFrame(columns=output_cols)
        excluded_zero_zones = 0
    grouped = (clean.groupby(['용도지역코드','용도지역명'],as_index=False)
               .agg(필지수=('PNU','nunique'),
                    교차구적면적합계_m2=('교차구적면적_m2','sum'),
                    최종보정면적합계_m2=('최종보정면적_m2','sum'))
               if not clean.empty else pd.DataFrame(columns=['용도지역코드','용도지역명','필지수','교차구적면적합계_m2','최종보정면적합계_m2']))
    pending = validations[validations['보정상태'] == '보정 보류'].copy()
    review = validations[validations['보정상태'] == '보정 완료·공간검토필요'].copy()
    passed = validations[validations['보정상태'].isin(['정상 보정', '보정 완료·공간검토필요'])]
    passed_count = len(passed)
    passed_target = float(passed['원본편입면적'].sum())
    passed_total = float(passed['최종보정합계_m2'].sum())
    if abs(passed_target-passed_total) > 10**(-precision)*1e-7:
        raise ArithmeticError('전체 정상 필지의 최종 면적 합계가 편입면적 합계와 다릅니다.')
    # Post-export invariant: even after zero rows are excluded, final Excel
    # zone totals must still match each adjusted PNU's source included area.
    exported_sums = {}
    for pnu, segment in clean.groupby('PNU', dropna=False):
        exported_sums[str(pnu).strip()] = sum(
            (Decimal(str(x)) for x in segment['최종보정면적_m2']), Decimal(0)
        )
    for _, parcel in passed.iterrows():
        pnu = str(parcel['PNU']).strip()
        expected = _as_decimal(parcel['원본편입면적'])
        actual = exported_sums.get(pnu, Decimal(0))
        if expected is None or abs(actual - expected) > Decimal('0.0000001'):
            raise ArithmeticError(f'{pnu}: 0㎡ 항목 제외 후에도 편입면적 합계가 일치해야 합니다. ({actual} != {expected})')
    total_pending = float(pd.to_numeric(pending['원본편입면적'],errors='coerce').fillna(0).sum())
    description = pd.DataFrame([
        ('기능','2-2 용도지역 교차분석과 편입면적 자동 보정 통합'),
        ('기준','PNU별 최종보정면적 합계 = 해당 필지 원본편입면적'),
        ('반올림 방식','PNU별 원본 편입면적 소수 자릿수에 맞춰 비례 배분 후 최대잔여법'),
        ('입력 면적 표시','정수/소수 한 자리/두 자리 혼재 가능. 엑셀에는 불필요한 소수 0 미표시'),
        ('필지별 보정 단위','정수 편입면적: 1㎡ / 소수 한 자리: 0.1㎡ / 소수 두 자리: 0.01㎡'),
        ('전체 대상 필지수',len(validations)),
        ('보정 완료 필지수',passed_count),
        ('최종 엑셀에서 제외한 0㎡ 용도지역 행수',excluded_zero_zones),
        ('최종 SHP에서 제외할 0㎡ 교차 피처수',len(excluded_zero_segments)),
        ('0㎡ 제외 기준','최종보정면적이 정확히 0㎡이고 보정 완료 상태인 항목만 최종 Excel/SHP에서 제외. 보류 행은 유지'),
        ('그중 공간검토 필요 필지수',len(review)),
        ('보정 보류 필지수',len(pending)),
        ('정상 필지 편입면적 합계(㎡)',passed_target),
        ('정상 필지 보정면적 합계(㎡)',passed_total),
        ('보류 필지 편입면적 합계(㎡) (참고)',total_pending),
        ('검증 주의','공간면적 차이는 보정하되 검토 필요로 따로 표시. PNU 중복·코드 오류·대장면적 초과 등은 보정 보류'),
        ('공간검토 의미','원인 미확정인 잠재적 누락·중복. 검토 필요 결과를 법적 확정면적으로 사용하지 마세요.'),
        ('유의사항','법정 산정 확정이 아니며 QGIS $area(타원체 설정)과 차이 가능'),
    ],columns=['항목','설명/결과'])
    visible_detail = raw[['분석행ID','PNU'] + source_extra_fields + ['지목','용도지역코드','용도지역명',
                          '원본편입면적','원본대장면적','교차구적면적_m2',
                          '최종보정면적_m2','보정차이_m2','적용보정단위_m2','보정상태','보류사유']]
    zero_audit = visible_detail.loc[zero_segment_mask].copy()
    if not zero_audit.empty:
        zero_audit['제외사유'] = '용도지역별 최종보정면적 0㎡ (최종 엑셀·SHP에서 제외)'
    else:
        zero_audit['제외사유'] = pd.Series(dtype='object')
    exported = _write_xlsx({
        '용도지역별_최종': clean[output_cols],
        '용도지역별_집계': grouped,
        '필지별_검증': validations,
        '보정상세': visible_detail,
        '0면적_제외내역': zero_audit,
        '보정보류_목록': pending,
        '공간검토필요_목록': review,
        '처리요약': description,
    }, precision=precision)
    return AdjustmentResult(
        excel_bytes=exported,
        adjusted_parcels=passed_count,
        pending_parcels=len(pending),
        total_parcels=len(validations),
        adjusted_area_m2=passed_total,
        pending_area_m2=total_pending,
        summary=grouped,
        preview=clean.head(100),
        pending=pending,
        review=review,
        review_parcels=len(review),
        validation=validations,
        detail=raw,
        excluded_zero_segments=len(excluded_zero_segments),
        excluded_zero_zones=excluded_zero_zones,
    )
