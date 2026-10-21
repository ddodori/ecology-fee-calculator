"""건축HUB 엑셀 두 형식에서 19자리 PNU를 생성하는 처리 엔진 (V1.2).

이 파일만으로 Streamlit 화면에 종속되지 않게 처리할 수 있습니다.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path
import csv
import re
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASIC_REQUIRED = ('시군구코드', '법정동코드', '대지구분코드', '번', '지')
CUSTOM_REQUIRED = ('시도', '시군구', '법정동', '대지구분', '번', '지')
APPENDED = ('법정동코드_10자리', 'PNU', '처리결과')


def _clean(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _numeric_part(value, width: int, allow_zero: bool = True) -> Optional[str]:
    """숫자 부분을 문자열로 보존하고, 앞의 0을 채웁니다.

    유효하지 않은 실수/문자/음수/길이 초과는 억지 변환하지 않습니다.
    """
    v = _clean(value)
    if not v or not re.fullmatch(r'\d+', v) or len(v) > width:
        return None
    n = int(v)
    if not allow_zero and n == 0:
        return None
    return str(n).zfill(width)


def read_codes(content: bytes) -> tuple[dict[str, str], set[str]]:
    """행정표준코드관리시스템 다운로드 TXT (CP949 또는 UTF-8)."""
    decoded = None
    for encoding in ('cp949', 'utf-8-sig', 'utf-8'):
        try:
            candidate = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if '법정동코드' in candidate.splitlines()[0]:
            decoded = candidate
            break
    if decoded is None:
        raise ValueError('법정동코드 TXT를 읽지 못했습니다. 공식 원본 TXT를 사용해 주세요.')
    data = csv.DictReader(decoded.splitlines(), delimiter='\t')
    codes: dict[str, str] = {}
    ambiguous: set[str] = set()
    for item in data:
        code = _clean(item.get('법정동코드'))
        name = ' '.join(_clean(item.get('법정동명')).split())
        state = _clean(item.get('폐지여부'))
        if not re.fullmatch(r'\d{10}', code) or not name or state != '존재':
            continue
        if name in codes and codes[name] != code:
            ambiguous.add(name)
        else:
            codes[name] = code
    for name in ambiguous:
        codes.pop(name, None)
    return codes, ambiguous


def _kind_from_basic(value) -> tuple[Optional[str], str]:
    """건축HUB 대지구분코드 0/1/2 => PNU 구분 1/2/제외."""
    code = _clean(value)
    if code == '0':
        return '1', '대지'
    if code == '1':
        return '2', '산'
    if code == '2':
        return None, '블록'
    return None, '확인 필요'


def _kind_from_custom(value) -> tuple[Optional[str], str]:
    kind = _clean(value)
    if kind == '대지':
        return '1', '대지'
    if kind == '산':
        return '2', '산'
    if kind == '블록':
        return None, '블록'
    return None, '확인 필요'


def _detect(headers: list[str]) -> str:
    cols = set(headers)
    if all(x in cols for x in BASIC_REQUIRED):
        return '기본개요'
    if all(x in cols for x in CUSTOM_REQUIRED):
        return '맞춤형 건축정보'
    basic_missing = ', '.join(x for x in BASIC_REQUIRED if x not in cols)
    custom_missing = ', '.join(x for x in CUSTOM_REQUIRED if x not in cols)
    raise ValueError(
        '지원하는 건축물대장 양식이 아닙니다.\n'
        f'기본개요에서 부족한 필드: {basic_missing}\n'
        f'맞춤형 건축정보에서 부족한 필드: {custom_missing}'
    )


def _prepare_output(headers: list[str]):
    wb = Workbook()
    ok = wb.active
    ok.title = 'PNU_생성결과'
    rejected = wb.create_sheet('PNU_제외목록')
    unique = wb.create_sheet('PNU_중복제거')
    summary = wb.create_sheet('처리요약')
    for sheet in (ok, rejected):
        sheet.append(headers + list(APPENDED))
    unique.append(['PNU', '법정동코드_10자리', '대지구분', '본번', '부번', '건축물대장_행수'])
    summary.append(['항목', '내용'])
    return wb, ok, rejected, unique, summary


def _style_output(wb, headers: list[str]):
    dark = PatternFill(fill_type='solid', fgColor='213A57')
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions if sheet.title != '처리요약' else 'A1:B1'
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.fill = dark
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(vertical='center')
        for col in range(1, sheet.max_column + 1):
            label = sheet.cell(1, col).value
            width = 28 if label in ('처리결과', '항목', '내용') else 23 if label in ('대지위치', '도로명대지위치') else 22 if label in ('PNU', '법정동코드_10자리') else 17
            sheet.column_dimensions[get_column_letter(col)].width = width
    for sheet in (wb['PNU_생성결과'], wb['PNU_제외목록'], wb['PNU_중복제거']):
        labels = [x.value for x in sheet[1]]
        text_cols = [i + 1 for i, name in enumerate(labels) if name in ('PNU', '법정동코드_10자리', '시군구코드', '법정동코드', '번', '지')]
        for cells in sheet.iter_rows(min_row=2):
            for c in text_cols:
                cells[c - 1].number_format = '@'


def process(excel_bytes: bytes, code_bytes: bytes | None = None):
    """return: (완성 엑셀 bytes, 통계 dict, 인식된 형식명)."""
    if not excel_bytes:
        raise ValueError('건축물대장 엑셀 파일을 올려주세요.')
    try:
        input_wb = load_workbook(BytesIO(excel_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError('엑셀 파일을 읽지 못했습니다. .xlsx 형식인지 확인해 주세요.') from exc

    try:
        sheet = input_wb.active
        # 일부 건축HUB 엑셀은 dimension이 A1만 기록돼서 범위를 재인식해야 합니다.
        if hasattr(sheet, 'reset_dimensions'):
            sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        first = next(rows, None)
        if first is None:
            raise ValueError('엑셀에 내용이 없습니다.')
        headers = [_clean(x) for x in first]
        # 건축HUB 기본개요 원본은 '구역코드명' 머리글이 실제로 2회 등장합니다.
        # 분석에 사용하는 필수 열만 중복 여부를 검사하고, 나머지 원본 열은 그대로 보존합니다.
        if any(headers.count(h) > 1 for h in set(BASIC_REQUIRED + CUSTOM_REQUIRED) if h in headers):
            raise ValueError('PNU 생성에 필요한 필드명이 중복되어 있습니다.')
        if any(x in headers for x in APPENDED):
            raise ValueError('이미 PNU 결과 필드가 포함된 파일입니다. 원본 건축물대장을 올려주세요.')
        source_type = _detect(headers)
        positions = {v: i for i, v in enumerate(headers)}

        if source_type == '맞춤형 건축정보':
            if code_bytes is None:
                code_bytes = Path(__file__).with_name('법정동코드_전체자료.txt').read_bytes()
            codes, ambiguous = read_codes(code_bytes)
        else:
            codes, ambiguous = {}, set()

        out, ok, bad, unique, summary = _prepare_output(headers)
        stats = Counter()
        counts_by_pnu = defaultdict(int)
        for data in rows:
            values = list(data[:len(headers)])
            values.extend([None] * (len(headers) - len(values)))
            if not any(_clean(v) for v in values):
                continue
            stats['전체'] += 1

            def get(field):
                return values[positions[field]]

            if source_type == '기본개요':
                city = _numeric_part(get('시군구코드'), 5, False)
                dong = _numeric_part(get('법정동코드'), 5, False)
                code10 = city + dong if city and dong else None
                kind_number, kind_name = _kind_from_basic(get('대지구분코드'))
                area_bad = None if code10 else '제외: 시군구/법정동코드 누락·오류'
            else:
                full_name = ' '.join(' '.join(_clean(get(x)) for x in ('시도', '시군구', '법정동')).split())
                code10 = codes.get(full_name)
                kind_number, kind_name = _kind_from_custom(get('대지구분'))
                if full_name in ambiguous:
                    area_bad = '제외: 법정동명 중복 코드'
                elif code10 is None:
                    area_bad = '제외: 법정동코드 미매칭'
                else:
                    area_bad = None

            bun = _numeric_part(get('번'), 4, False)
            ji = _numeric_part(get('지'), 4, True)
            if kind_name == '블록':
                reason = '제외: 블록'
            elif kind_number is None:
                reason = '제외: 대지구분 확인'
            elif area_bad:
                reason = area_bad
            elif bun is None:
                reason = '제외: 본번 누락/오류'
            elif ji is None:
                reason = '제외: 부번 누락/오류'
            else:
                reason = '정상'
            pnu = (code10 + kind_number + bun + ji) if reason == '정상' else None
            result = values + [code10, pnu, reason]

            if reason == '정상':
                ok.append(result)
                stats['정상'] += 1
                counts_by_pnu[pnu] += 1
            else:
                bad.append(result)
                stats[reason] += 1

        # 중복 제거는 PNU 결과만 독립 목록으로 만듭니다. 원본 행은 유지합니다.
        for pnu, n in counts_by_pnu.items():
            kind_name = '산' if pnu[10] == '2' else '대지'
            unique.append([pnu, pnu[:10], kind_name, pnu[11:15], pnu[15:19], n])

        stats['고유 PNU'] = len(counts_by_pnu)
        stats['중복 PNU 행 수'] = stats['정상'] - stats['고유 PNU']
        stats['제외'] = stats['전체'] - stats['정상']
        summary.append(['인식된 입력 형식', source_type])
        summary.append(['전체 데이터', stats['전체']])
        summary.append(['정상 PNU 행', stats['정상']])
        summary.append(['고유 PNU 수', stats['고유 PNU']])
        summary.append(['중복 PNU 행 수', stats['중복 PNU 행 수']])
        summary.append(['제외 행', stats['제외']])
        for reason in sorted(stats):
            if reason.startswith('제외:'):
                summary.append([reason, stats[reason]])
        if source_type == '기본개요':
            summary.append(['법정동코드 출처', '엑셀 내 시군구코드(5자리) + 법정동코드(5자리)'])
        else:
            summary.append(['법정동코드 출처', '법정동코드 TXT — 폐지여부=존재'])
        summary.append(['변환 규칙', '건축HUB 대지구분 0→PNU 1(대지), 1→2(산), 2→블록 제외'])
        summary.append(['주의', '제외 행은 건축물이 없다는 뜻이 아님. 필지 및 코드 기준시점 확인 필요'])

        _style_output(out, headers)
        result_file = BytesIO()
        out.save(result_file)
        return result_file.getvalue(), dict(stats), source_type
    finally:
        input_wb.close()
