"""프로그램이 정상 설치됐는지 확인하는 간단한 자체 테스트."""
from io import BytesIO
from openpyxl import Workbook, load_workbook
from pnu_core import process

wb = Workbook()
ws = wb.active
ws.append(['시군구코드','법정동코드','대지구분코드','번','지'])
ws.append(['41595','10100',0,4,17])
ws.append(['41595','10100',1,82,0])
ws.append(['41595','10100',2,0,0])
stream = BytesIO()
wb.save(stream)
output, counts, kind = process(stream.getvalue())
result = load_workbook(BytesIO(output), read_only=True)
pnus = [row[-2] for row in list(result['PNU_생성결과'].values)[1:]]
assert kind == '기본개요'
assert pnus == ['4159510100100040017', '4159510100200820000']
assert counts['제외: 블록'] == 1
print('자체 테스트 성공: 대지/산 PNU 생성 및 블록 제외 정상')
