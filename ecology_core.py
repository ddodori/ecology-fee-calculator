"""2-4 생태자연도 교차 및 면적 보정 (V1.4.18).

Input: 2-3 제외후_잔여 SHP ZIP, EcoBank 생태자연도 SHP ZIP.
The complement of grades 1, 2 and 별도관리지역 is grade 3, even when
no grade-3 polygon exists in the ecological input.
The target of each remaining zoning fragment is its ALREADY adjusted area
('보정면적'), not the projected geometry area. Allocation uses largest remainders
at the decimal precision of that fragment, preserving exact totals.
The class priority 별도 > 1 > 2 is only a deterministic treatment of overlaps;
any cross-class overlaps are explicitly flagged for manual review.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_FLOOR
from io import BytesIO
from collections import defaultdict
from typing import Iterable
import math
import re

import geopandas as gpd
import pandas as pd
from shapely.ops import unary_union

from zone_core import read_spatial, WORK_CRS, _polygon_only
from shp_export import make_shp_zip

CLASSES = ('1등급', '2등급', '3등급', '별도관리지역')
PRIORITY = ('별도관리지역', '1등급', '2등급')


@dataclass
class EcologyResult:
    excel_bytes: bytes
    shp_zip: bytes | None
    summary: pd.DataFrame
    details: pd.DataFrame
    validation: pd.DataFrame
    total_m2: Decimal
    class_totals: dict[str, Decimal]
    overlap_fragments: int
    overlap_area_m2: float
    feature_count: int


def _jibeon(pnu: str) -> str:
    key = str(pnu).strip()
    if len(key) != 19 or not key.isdigit() or key[10] not in ('1', '2'):
        return ''
    main, sub = int(key[11:15]), int(key[15:19])
    if not main:
        return ''
    return ('산' if key[10] == '2' else '') + str(main) + (f'-{sub}' if sub else '')


def _clean_val(v: object) -> str:
    if v is None or pd.isna(v):
        return ''
    return str(v).strip()


def default_grade_values(values: Iterable[object]) -> dict[str, list[str]]:
    """Only map meaning that is explicitly readable; never guess code 0 or B."""
    out = {'1등급': [], '2등급': [], '별도관리지역': []}
    for value in values:
        s = _clean_val(value)
        upper = re.sub(r'\s+', '', s).upper()
        if upper in ('1', '1.0', '1등급', 'Ⅰ', 'I', 'GRADE1'):
            out['1등급'].append(s)
        elif upper in ('2', '2.0', '2등급', 'Ⅱ', 'II', 'GRADE2'):
            out['2등급'].append(s)
        elif any(token in upper for token in ('별도', '별도관리', 'SPECIAL')):
            out['별도관리지역'].append(s)
    return out


def _to_area(v: object, label: str) -> Decimal:
    try:
        if v is None or pd.isna(v):
            raise ValueError()
        val = Decimal(str(v).replace(',', '').strip())
        if not val.is_finite() or val < 0:
            raise ValueError()
        return val
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f'{label}: 보정면적이 비었거나 유효한 0 이상 숫자가 아닙니다: {v!r}')


def _unit(v: Decimal) -> Decimal:
    """Integer 1m², one decimal 0.1m², two decimals 0.01m²."""
    if v == v.to_integral_value():
        return Decimal('1')
    if v * 10 == (v * 10).to_integral_value():
        return Decimal('0.1')
    if v * 100 == (v * 100).to_integral_value():
        return Decimal('0.01')
    raise ValueError(f'편입 보정면적 {v}㎡가 소수 셋째 자리 이상입니다. 임의 반올림하지 않습니다.')


def _allocate(target: Decimal, raw: dict[str, float]) -> dict[str, Decimal]:
    """Hamilton largest-remainder allocation with stable deterministic ties."""
    quantum = _unit(target)
    base = sum(raw.values())
    if base <= 0 or not math.isfinite(base):
        raise ValueError('잔여면적 도형의 면적이 0이거나 유효하지 않습니다.')
    units = int(target / quantum)
    ideal = [Decimal(units) * Decimal(str(max(0.0, raw[k]))) / Decimal(str(base)) for k in CLASSES]
    fixed = [int(x.to_integral_value(rounding=ROUND_FLOOR)) for x in ideal]
    remainder = units - sum(fixed)
    ranked = sorted(range(len(CLASSES)), key=lambda idx: (-(ideal[idx] - fixed[idx]), idx))
    for idx in ranked[:remainder]:
        fixed[idx] += 1
    values = {key: quantum * fixed[i] for i, key in enumerate(CLASSES)}
    assert sum(values.values(), Decimal(0)) == target
    return values


def _clean_shape(geom):
    poly = _polygon_only(geom)
    if poly is None or poly.is_empty:
        return None
    return poly


def _find_col(frame: pd.DataFrame, options: tuple[str, ...]) -> str | None:
    normalized = {re.sub(r'[^a-zA-Z0-9가-힣]', '', str(c)).casefold(): str(c) for c in frame.columns}
    for o in options:
        col = normalized.get(re.sub(r'[^a-zA-Z0-9가-힣]', '', o).casefold())
        if col:
            return col
    return None


def _make_excel(detail: pd.DataFrame, validation: pd.DataFrame,
                totals: dict[str, Decimal], overlap_fragments: int, overlap_area: float) -> bytes:
    """Reviewer-oriented Excel: two sheets, no redundant diagnostic columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook()
    summary = book.active
    summary.title = '생태자연도_집계'
    data = book.create_sheet('필지별_생태등급')
    summary.merge_cells('A1:D1')
    summary['A1'] = '생태자연도 등급별 최종면적'
    summary.merge_cells('A2:D2')
    summary['A2'] = '1·2등급·별도관리지역을 제외한 나머지는 3등급 | 제외 후 잔여 보정면적 기준'
    summary.append([])
    summary.append(['등급', '최종면적(㎡)', '비율(%)', '설명'])
    total = sum(totals.values(), Decimal(0))
    for idx, grade in enumerate(CLASSES, 5):
        summary.cell(idx, 1, grade)
        summary.cell(idx, 2, float(totals[grade]))
        summary.cell(idx, 3, float(totals[grade] / total * 100) if total else 0)
        summary.cell(idx, 4, '잔여공간에서 자동 생성' if grade == '3등급' else '생태자연도 SHP와 교차')
    summary['A9'] = '합계'
    summary['B9'] = '=SUM(B5:B8)'
    summary['C9'] = '=SUM(C5:C8)'
    summary['A11'] = '원본 잔여면적 합계'
    summary['B11'] = float(total)
    summary['A12'] = '보정합계 차이'
    summary['B12'] = '=B9-B11'
    summary['A14'] = '등급 경계 중복 검토'
    summary['B14'] = f'{overlap_fragments}개 잔여 구간 / 약 {overlap_area:,.2f}㎡'
    summary.merge_cells('A15:D15')
    summary['A15'] = ('서로 다른 등급의 도형이 중첩되면 별도관리 > 1등급 > 2등급을 임시 우선 적용. '
                      '실제 경계·원본 자료 확인 필요.')
    summary['A15'].alignment = Alignment(wrap_text=True, vertical='center')
    summary.row_dimensions[15].height = 36

    headers = ['연번', 'PNU', '지번', '지목', '용도지역', '잔여면적(㎡)',
               '1등급(㎡)', '2등급(㎡)', '3등급(㎡)', '별도관리(㎡)', '등급합계(㎡)', '차이(㎡)', '검토']
    data.append(headers)
    for idx, row in enumerate(detail.to_dict('records'), 2):
        data.append([idx-1, row['PNU'], row['지번'], row['지목'], row['용도지역'], row['잔여면적(㎡)'],
                     row['1등급(㎡)'], row['2등급(㎡)'], row['3등급(㎡)'], row['별도관리(㎡)'],
                     None, None, row['검토']])
        data.cell(idx, 11, f'=SUM(G{idx}:J{idx})')
        data.cell(idx, 12, f'=K{idx}-F{idx}')
        data.cell(idx, 2).number_format = '@'
    data.freeze_panes = 'F2'
    data.auto_filter.ref = f'A1:M{max(2,1+len(detail))}'
    data.sheet_view.showGridLines = False
    summary.sheet_view.showGridLines = False
    for sheet, head_row, col_count in ((summary, 4, 4), (data, 1, 13)):
        sheet.row_dimensions[head_row].height = 27
        for cell in sheet[head_row]:
            if cell.column <= col_count:
                cell.fill = PatternFill('solid', fgColor='244F64')
                cell.font = Font(bold=True, color='FFFFFF', size=10)
                cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.freeze_panes = sheet.freeze_panes or f'A{head_row+1}'
    for col, width in zip('ABCD', (26, 24, 16, 80)):
        summary.column_dimensions[col].width = 46 if col == 'D' else width
    for col, width in zip('ABCDEFGHIJKLM', (8,24,16,10,24,18,16,16,16,18,18,16,31)):
        data.column_dimensions[col].width = width
    def number_display(value):
        try:
            unit = _unit(Decimal(str(value)))
            return '#,##0' if unit == 1 else '#,##0.0' if unit == Decimal('0.1') else '#,##0.00'
        except (ValueError, InvalidOperation):
            return '#,##0.00'
    for row_idx, cls in enumerate(CLASSES, 5):
        summary.cell(row_idx, 2).number_format = number_display(totals[cls])
        summary.cell(row_idx, 3).number_format = '0.00'
    for cell_ref in ('B9','B11','B12'):
        summary[cell_ref].number_format = number_display(total)
    summary['C9'].number_format = '0.00'
    for idx, record in enumerate(detail.to_dict('records'),2):
        for col_idx in range(6, 11):
            value = data.cell(idx,col_idx).value
            data.cell(idx,col_idx).number_format = number_display(value)
        for col_idx in (11,12):
            data.cell(idx,col_idx).number_format = number_display(record['잔여면적(㎡)'])
    for sheet in (summary, data):
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output=BytesIO()
    book.save(output)
    return output.getvalue()


def analyze_ecology(
    remaining_file: tuple[str, bytes],
    ecology_files: list[tuple[str, bytes]],
    *, grade_field: str,
    special_values: Iterable[str],
    grade1_values: Iterable[str],
    grade2_values: Iterable[str],
    pnu_field: str = 'PNU',
    area_field: str = '보정면적',
    remaining_encoding: str = 'auto',
    ecology_encoding: str = 'auto',
) -> EcologyResult:
    if not ecology_files:
        raise ValueError('생태자연도 SHP ZIP을 한 개 이상 업로드하세요.')
    remaining = read_spatial(remaining_file[1], remaining_file[0], encoding=remaining_encoding)
    if pnu_field not in remaining or area_field not in remaining:
        raise ValueError(f'잔여 SHP에 {pnu_field}, {area_field} 필드가 필요합니다.')
    remaining = remaining.to_crs(WORK_CRS)
    remaining = remaining.loc[~remaining.geometry.is_empty].copy()
    if remaining.empty:
        raise ValueError('잔여 도형이 없습니다.')
    sources = []
    for name, payload in ecology_files:
        eco = read_spatial(payload, name, encoding=ecology_encoding)
        if grade_field not in eco:
            raise ValueError(f'{name}: 생태자연도 등급 필드 {grade_field!r}가 없습니다.')
        sources.append(eco[[grade_field, eco.geometry.name]].to_crs(WORK_CRS).rename_geometry('geometry') if eco.geometry.name != 'geometry' else eco[[grade_field, 'geometry']].to_crs(WORK_CRS))
    ecology = gpd.GeoDataFrame(pd.concat(sources, ignore_index=True), geometry='geometry', crs=WORK_CRS)
    sets = {
        '1등급': set(map(str, grade1_values)),
        '2등급': set(map(str, grade2_values)),
        '별도관리지역': set(map(str, special_values)),
    }
    for left in ('1등급', '2등급', '별도관리지역'):
        for right in ('1등급', '2등급', '별도관리지역'):
            if left < right and sets[left] & sets[right]:
                raise ValueError(f'등급값이 중복 지정됐습니다: {left}/{right} {sets[left]&sets[right]}')
    classified = {}
    values = ecology[grade_field].map(_clean_val)
    for cls in PRIORITY:
        cls_frame = ecology.loc[values.isin(sets[cls])].copy()
        classified[cls] = cls_frame
    indexes = {cls: frame.sindex if not frame.empty else None for cls, frame in classified.items()}
    fields = {
        'pnu': pnu_field,
        'use': _find_col(remaining, ('지목','1_지목')),
        'zone': _find_col(remaining, ('용도지역', '용도지역명', '분석용도')),
    }
    totals = {cls: Decimal(0) for cls in CLASSES}
    records = []
    spatial_rows = []
    validations = []
    overlaps = 0
    overlap_area = 0.0
    for _, origin in remaining.iterrows():
        pnu = _clean_val(origin[pnu_field])
        if not pnu:
            raise ValueError('잔여 SHP에 PNU가 비어 있는 행이 있습니다.')
        target = _to_area(origin[area_field], f'PNU {pnu}')
        if target <= 0:
            continue
        shape = _clean_shape(origin.geometry)
        if shape is None or shape.area <= 0:
            raise ValueError(f'PNU {pnu}: 잔여 도형이 유효하지 않습니다.')
        assigned_geom = {}
        occupied = None
        warned = False
        for cls in PRIORITY:
            class_frame = classified[cls]
            sindex = indexes[cls]
            intersected = []
            if sindex is not None:
                for index in sindex.query(shape, predicate='intersects'):
                    item = class_frame.geometry.iloc[int(index)]
                    part = _clean_shape(shape.intersection(item))
                    if part is not None and part.area > 0:
                        intersected.append(part)
            mask = _clean_shape(unary_union(intersected)) if intersected else None
            if mask is None:
                assigned_geom[cls] = None
                continue
            if occupied is not None:
                shared = mask.intersection(occupied).area
                if shared > 1e-5:
                    warned = True
                    overlap_area += shared
                mask = _clean_shape(mask.difference(occupied))
            assigned_geom[cls] = mask
            if mask is not None:
                occupied = _clean_shape(unary_union([occupied, mask])) if occupied is not None else mask
        grade3 = _clean_shape(shape.difference(occupied)) if occupied is not None else shape
        assigned_geom['3등급'] = grade3
        raw = {cls: max(0.0, assigned_geom[cls].area) if assigned_geom.get(cls) is not None else 0.0 for cls in CLASSES}
        # An imperfect geometry repair may leave tiny uncovered/over-covered parts,
        # so check that all mutually disjoint pieces sum to the original polygon.
        if abs(sum(raw.values()) - shape.area) > max(0.001, shape.area * 1e-7):
            raise ArithmeticError(f'PNU {pnu}: 공간 분할 결과의 면적 합계가 잔여 도형과 일치하지 않습니다.')
        allocated = _allocate(target, raw)
        if warned:
            overlaps += 1
        for cls in CLASSES:
            totals[cls] += allocated[cls]
            if allocated[cls] > 0 and assigned_geom[cls] is not None:
                spatial_rows.append({
                    'PNU': pnu, '지목': _clean_val(origin[fields['use']]) if fields['use'] else '',
                    '용도지역': _clean_val(origin[fields['zone']]) if fields['zone'] else '',
                    '생태등급': cls, '등급면적': float(allocated[cls]),
                    'geometry': assigned_geom[cls],
                })
        record = {'PNU':pnu, '지번':_jibeon(pnu), '지목':_clean_val(origin[fields['use']]) if fields['use'] else '',
                  '용도지역':_clean_val(origin[fields['zone']]) if fields['zone'] else '',
                  '잔여면적(㎡)':float(target), '검토':'등급 도형 중첩' if warned else ''}
        for cls in CLASSES:
            record[f'{cls if cls != "별도관리지역" else "별도관리"}(㎡)'] = float(allocated[cls])
        records.append(record)
        validations.append({'PNU':pnu,'기준잔여면적':target,'등급면적합계':sum(allocated.values(),Decimal(0)),
                            '일치':sum(allocated.values(),Decimal(0)) == target})
    if not records:
        raise ValueError('보정면적이 0㎡보다 큰 잔여 도형이 없습니다.')
    check = pd.DataFrame(validations)
    if not check['일치'].all():
        raise ArithmeticError('등급 보정면적 합계가 잔여면적과 일치하지 않습니다.')
    detail = pd.DataFrame(records)
    summary = pd.DataFrame([{'생태등급':cls, '보정면적_m2':float(totals[cls])} for cls in CLASSES])
    spatial = gpd.GeoDataFrame(spatial_rows, geometry='geometry', crs=WORK_CRS)
    shp_bytes = make_shp_zip(spatial, base_name='생태자연도_등급교차결과') if not spatial.empty else None
    excel_bytes = _make_excel(detail, check, totals, overlaps, overlap_area)
    return EcologyResult(
        excel_bytes=excel_bytes, shp_zip=shp_bytes, summary=summary, details=detail,
        validation=check, total_m2=sum((v['기준잔여면적'] for v in validations),Decimal(0)),
        class_totals=totals, overlap_fragments=overlaps, overlap_area_m2=overlap_area,
        feature_count=len(spatial),
    )
